from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Optional


class CsvImportService:
    EXPECTED_COLUMNS = {
        "borrower_name", "borrower_group", "amount", "giving_date",
        "depositor_name", "depositor_group", "due_date", "reference_id",
    }

    DATE_FORMATS = ["%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"]

    def parse(self, file_path: str) -> list[dict]:
        """Parse CSV or XLSX file into list of dicts with normalized keys."""
        path = Path(file_path)
        if path.suffix.lower() == ".xlsx":
            return self._parse_xlsx(path)
        return self._parse_csv(path)

    def _parse_csv(self, path: Path) -> list[dict]:
        records = []
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                normalized = self._normalize_row(row)
                records.append(normalized)
        return records

    def _parse_xlsx(self, path: Path) -> list[dict]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError("openpyxl is required for XLSX import")

        wb = load_workbook(path, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h).lower().strip() if h else "" for h in rows[0]]
        records = []
        for row in rows[1:]:
            row_dict = {}
            for i, header in enumerate(headers):
                if i < len(row):
                    row_dict[header] = row[i]
            normalized = self._normalize_row(row_dict)
            records.append(normalized)
        wb.close()
        return records

    def _normalize_row(self, row: dict) -> dict:
        """Normalize column names and parse values."""
        normalized: dict = {}
        cleaned = {k.lower().strip(): v for k, v in row.items()}

        normalized["borrower_name"] = self._clean_str(cleaned.get("borrower_name"))
        normalized["borrower_group"] = self._clean_str(cleaned.get("borrower_group"))
        normalized["depositor_name"] = self._clean_str(cleaned.get("depositor_name"))
        normalized["depositor_group"] = self._clean_str(cleaned.get("depositor_group"))
        normalized["amount"] = self._parse_int(cleaned.get("amount"))
        normalized["giving_date"] = self._parse_date(cleaned.get("giving_date"))
        normalized["due_date"] = self._parse_date(cleaned.get("due_date"))

        ref_id = cleaned.get("reference_id")
        normalized["reference_id"] = ref_id.strip() if ref_id and str(ref_id).strip() else None

        return normalized

    @staticmethod
    def _clean_str(value: Optional[object]) -> Optional[str]:
        if value is None:
            return None
        s = str(value).strip()
        return s.lower() if s else None

    def _parse_date(self, value: Optional[object]) -> Optional[date]:
        if value is None:
            return None
        if isinstance(value, date):
            return value
        if isinstance(value, datetime):
            return value.date()
        s = str(value).strip()
        if not s:
            return None
        for fmt in self.DATE_FORMATS:
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _parse_int(value: Optional[object]) -> Optional[int]:
        if value is None:
            return None
        try:
            return int(float(str(value).strip()))
        except (ValueError, TypeError):
            return None
